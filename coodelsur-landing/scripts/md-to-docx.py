"""Convierte WITME-API-PARA-INTEGRADOR.md a Word (.docx)."""
from __future__ import annotations

import re
import sys
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt


def add_inline_runs(paragraph, node) -> None:
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            paragraph.add_run(text)
        return

    if not isinstance(node, Tag):
        return

    if node.name == "strong":
        run = paragraph.add_run(node.get_text())
        run.bold = True
    elif node.name == "em":
        run = paragraph.add_run(node.get_text())
        run.italic = True
    elif node.name == "code":
        run = paragraph.add_run(node.get_text())
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    elif node.name == "a":
        run = paragraph.add_run(node.get_text())
        run.underline = True
    else:
        for child in node.children:
            add_inline_runs(paragraph, child)


def add_table(doc: Document, table_tag: Tag) -> None:
    rows = table_tag.find_all("tr")
    if not rows:
        return

    cols = max(len(row.find_all(["th", "td"])) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"

    for r, row in enumerate(rows):
        cells = row.find_all(["th", "td"])
        for c, cell in enumerate(cells):
            table.rows[r].cells[c].text = cell.get_text(strip=True)


def add_codeblock(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text.strip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    paragraph.paragraph_format.left_indent = Pt(12)


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    md_content = md_path.read_text(encoding="utf-8")
    html = markdown.markdown(
        md_content,
        extensions=["tables", "fenced_code", "sane_lists", "nl2br"],
    )
    soup = BeautifulSoup(html, "html.parser")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for element in soup.children:
        if not isinstance(element, Tag):
            continue

        if element.name == "h1":
            doc.add_heading(element.get_text(strip=True), level=0)
        elif element.name == "h2":
            doc.add_heading(element.get_text(strip=True), level=1)
        elif element.name == "h3":
            doc.add_heading(element.get_text(strip=True), level=2)
        elif element.name == "h4":
            doc.add_heading(element.get_text(strip=True), level=3)
        elif element.name == "p":
            paragraph = doc.add_paragraph()
            for child in element.children:
                add_inline_runs(paragraph, child)
        elif element.name == "ul":
            for li in element.find_all("li", recursive=False):
                paragraph = doc.add_paragraph(style="List Bullet")
                for child in li.children:
                    add_inline_runs(paragraph, child)
        elif element.name == "ol":
            for li in element.find_all("li", recursive=False):
                paragraph = doc.add_paragraph(style="List Number")
                for child in li.children:
                    add_inline_runs(paragraph, child)
        elif element.name == "table":
            add_table(doc, element)
            doc.add_paragraph()
        elif element.name == "pre":
            code = element.get_text()
            add_codeblock(doc, code)
        elif element.name == "hr":
            doc.add_paragraph("—" * 40).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        elif element.name == "blockquote":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(element.get_text(strip=True))
            run.italic = True

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1]
    md_file = root / "docs" / "WITME-API-PARA-INTEGRADOR.md"
    out_file = root / "docs" / "WITME-API-PARA-INTEGRADOR.docx"
    desktop = Path.home() / "Desktop" / "WITME-API-PARA-INTEGRADOR.docx"

    if len(sys.argv) > 1:
        out_file = Path(sys.argv[1])

    convert_markdown_to_docx(md_file, out_file)
    convert_markdown_to_docx(md_file, desktop)
    print(f"Generado: {out_file}")
    print(f"Generado: {desktop}")
